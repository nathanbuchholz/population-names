"""Per-source converted-tier transforms.

Each transform reads from archived/{source_id}/ and writes normalized CSVs
to converted/{source_id}/. If a source has a registered transform, it replaces
the generic extraction logic in FileStore.promote_to_converted.

Schema versions: Each source defines ``schema_versions`` in sources.yml with
``date_range: [min_year, max_year]`` (null = unbounded). The transform extracts
a date from the filename via ``filename_date_re`` and resolves which version
config to use per file.

Country-specific transforms live in submodules:
  - gb_england_wales: ONS forenames, Wales forenames, England surnames
  - gb_scotland: NRS forenames, NRS surnames
  - gb_northern_ireland: NISRA forenames
  - ie: CSO forenames (boys/girls), CSO surnames
  - us: SSA forenames, Census surnames
"""

import csv
import logging
import re
from pathlib import Path

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Registry
# ---------------------------------------------------------------------------

_TRANSFORMS: dict[str, callable] = {}


def register(source_id: str):
    """Decorator to register a transform function for a source."""

    def wrapper(fn):
        _TRANSFORMS[source_id] = fn
        return fn

    return wrapper


def get_transform(source_id: str):
    """Return the transform function for a source, or None."""
    return _TRANSFORMS.get(source_id)


# ---------------------------------------------------------------------------
# Schema version resolution
# ---------------------------------------------------------------------------


def resolve_schema_version(filename: str, cfg: dict) -> dict:
    """Extract date from filename and return matching schema version config.

    Uses ``store.filename_date_re`` to extract a year from the filename,
    then matches against ``schema_versions[*].date_range``.

    Returns the matching version dict, or ``{}`` if no versions are defined.
    """
    versions = cfg.get("schema_versions", [])
    if not versions:
        return {}

    date_re = cfg.get("store", {}).get("filename_date_re")
    if not date_re:
        return versions[0]

    match = re.search(date_re, filename)
    if not match:
        logger.warning(
            "No date match in %r using pattern %r, defaulting to first version",
            filename,
            date_re,
        )
        return versions[0]

    file_date = int(match.group(1))

    for version in versions:
        dr = version.get("date_range", [None, None])
        low, high = dr[0], dr[1]
        if (low is None or file_date >= low) and (high is None or file_date <= high):
            return version

    logger.warning(
        "No schema version matched date %d for %r, defaulting to last version",
        file_date,
        filename,
    )
    return versions[-1]


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def write_csv(path: Path, header: list[str], rows: list[list]) -> Path:
    """Write rows to a CSV file with the given header."""
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(header)
        writer.writerows(rows)
    logger.info("Converted: wrote %s (%d rows)", path.name, len(rows))
    return path


def read_xlsx_rows(xlsx_path: Path, sheet_name: str, skip_rows: int = 0):
    """Read rows from an XLSX sheet, skipping metadata rows."""
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        logger.warning("Sheet %r not found in %s", sheet_name, xlsx_path.name)
        wb.close()
        return []
    ws = wb[sheet_name]
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < skip_rows:
            continue
        rows.append(list(row))
    wb.close()
    return rows


# ---------------------------------------------------------------------------
# Import submodules to trigger @register decorators
# ---------------------------------------------------------------------------

from pipeline.transforms import (  # noqa: E402
    gb_england_wales,  # noqa: F401
    gb_northern_ireland,  # noqa: F401
    gb_scotland,  # noqa: F401
    ie,  # noqa: F401
    us,  # noqa: F401
)
