"""Tests for pipeline transform functions and helpers."""

import csv

from pipeline.transforms import read_xlsx_rows, resolve_schema_version, write_csv

# ---------------------------------------------------------------------------
# Helper tests
# ---------------------------------------------------------------------------


class TestResolveSchemaVersion:
    def test_no_versions(self):
        cfg = {"schema_versions": []}
        assert resolve_schema_version("file.csv", cfg) == {}

    def test_matches_date(self):
        cfg = {
            "store": {"filename_date_re": r"(\d{4})"},
            "schema_versions": [
                {"date_range": [2000, 2009], "label": "v1"},
                {"date_range": [2010, 2019], "label": "v2"},
            ],
        }
        result = resolve_schema_version("data_2015.csv", cfg)
        assert result["label"] == "v2"

    def test_no_date_re(self):
        cfg = {
            "store": {},
            "schema_versions": [{"label": "first"}, {"label": "second"}],
        }
        result = resolve_schema_version("anything.csv", cfg)
        assert result["label"] == "first"

    def test_no_match_falls_back(self):
        cfg = {
            "store": {"filename_date_re": r"(\d{4})"},
            "schema_versions": [
                {"date_range": [2000, 2009], "label": "v1"},
                {"date_range": [2010, 2019], "label": "v2"},
            ],
        }
        # Year 2025 doesn't match any range
        result = resolve_schema_version("data_2025.csv", cfg)
        assert result["label"] == "v2"  # falls back to last


class TestWriteCsv:
    def test_writes_header_and_data(self, tmp_path):
        out = tmp_path / "out.csv"
        write_csv(out, ["a", "b"], [[1, 2], [3, 4]])
        lines = out.read_text().strip().split("\n")
        assert lines[0] == "a,b"
        assert lines[1] == "1,2"
        assert lines[2] == "3,4"


class TestReadXlsxRows:
    def test_reads_rows(self, tmp_path):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["header1", "header2"])
        ws.append(["a", 1])
        ws.append(["b", 2])
        path = tmp_path / "test.xlsx"
        wb.save(path)
        wb.close()

        rows = read_xlsx_rows(path, "Sheet1")
        assert len(rows) == 3
        assert rows[0] == ["header1", "header2"]
        assert rows[1] == ["a", 1]

    def test_skip_rows(self, tmp_path):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["metadata"])
        ws.append(["more metadata"])
        ws.append(["name", "count"])
        ws.append(["Alice", 10])
        path = tmp_path / "test.xlsx"
        wb.save(path)
        wb.close()

        rows = read_xlsx_rows(path, "Data", skip_rows=2)
        assert len(rows) == 2
        assert rows[0] == ["name", "count"]

    def test_missing_sheet(self, tmp_path):
        import openpyxl

        wb = openpyxl.Workbook()
        path = tmp_path / "test.xlsx"
        wb.save(path)
        wb.close()

        rows = read_xlsx_rows(path, "NonExistent")
        assert rows == []


# ---------------------------------------------------------------------------
# US transforms
# ---------------------------------------------------------------------------


class TestSSAForenames:
    def _cfg(self):
        return {
            "store": {"filename_date_re": r"yob(\d{4})"},
            "schema_versions": [{"date_range": [None, None]}],
        }

    def test_basic(self, tmp_path):
        from pipeline.transforms.us import transform_ssa_forenames

        archived = tmp_path / "archived"
        converted = tmp_path / "converted"
        archived.mkdir()
        converted.mkdir()

        (archived / "yob2020.txt").write_text("Mary,F,1000\nJohn,M,900\n")

        results = transform_ssa_forenames(archived, converted, self._cfg())
        assert len(results) == 1

        with open(results[0]) as f:
            reader = csv.DictReader(f)
            rows = list(reader)
        assert len(rows) == 2
        assert rows[0] == {"name": "Mary", "sex": "F", "count": "1000", "year": "2020"}
        assert rows[1] == {"name": "John", "sex": "M", "count": "900", "year": "2020"}

    def test_skips_non_txt(self, tmp_path):
        from pipeline.transforms.us import transform_ssa_forenames

        archived = tmp_path / "archived"
        converted = tmp_path / "converted"
        archived.mkdir()
        converted.mkdir()

        (archived / "readme.pdf").write_text("not a data file")

        results = transform_ssa_forenames(archived, converted, self._cfg())
        assert results == []

    def test_skips_bad_filename(self, tmp_path):
        from pipeline.transforms.us import transform_ssa_forenames

        archived = tmp_path / "archived"
        converted = tmp_path / "converted"
        archived.mkdir()
        converted.mkdir()

        (archived / "notes.txt").write_text("Mary,F,1000\n")

        results = transform_ssa_forenames(archived, converted, self._cfg())
        assert results == []


class TestCensusSurnames:
    def test_basic(self, tmp_path):
        from pipeline.transforms.us import transform_census_surnames

        archived = tmp_path / "archived"
        converted = tmp_path / "converted"
        archived.mkdir()
        converted.mkdir()

        (archived / "surnames.csv").write_text("name,count\nSMITH,100\njohnson,200\n")

        col_map = {"name": "name", "count": "count"}
        cfg = {
            "store": {},
            "schema_versions": [
                {
                    "date_range": [None, None],
                    "column_map": col_map,
                    "default_year": 2010,
                }
            ],
        }
        results = transform_census_surnames(archived, converted, cfg)
        assert len(results) == 1

        with open(results[0]) as f:
            reader = csv.DictReader(f)
            rows = list(reader)
        assert rows[0]["name"] == "Smith"  # title-cased
        assert rows[1]["name"] == "Johnson"

    def test_default_year(self, tmp_path):
        from pipeline.transforms.us import transform_census_surnames

        archived = tmp_path / "archived"
        converted = tmp_path / "converted"
        archived.mkdir()
        converted.mkdir()

        (archived / "data.csv").write_text("name,count\nDoe,50\n")

        col_map = {"name": "name", "count": "count"}
        cfg = {
            "store": {},
            "schema_versions": [
                {
                    "date_range": [None, None],
                    "column_map": col_map,
                    "default_year": 2000,
                }
            ],
        }
        results = transform_census_surnames(archived, converted, cfg)
        with open(results[0]) as f:
            rows = list(csv.DictReader(f))
        assert rows[0]["year"] == "2000"


# ---------------------------------------------------------------------------
# Scotland (NRS) transforms
# ---------------------------------------------------------------------------


class TestNRSForenames:
    def test_basic(self, tmp_path):
        from pipeline.transforms.gb_scotland import transform_nrs_forenames

        archived = tmp_path / "archived"
        converted = tmp_path / "converted"
        archived.mkdir()
        converted.mkdir()

        (archived / "nrs.csv").write_text(
            "Year,Sex,Name,Number,Rank\n2020,Boy,Jack,500,1\n2020,Girl,Emily,450,2\n"
        )

        col_map = {
            "name": "Name",
            "sex": "Sex",
            "count": "Number",
            "rank": "Rank",
            "year": "Year",
        }
        cfg = {
            "store": {},
            "schema_versions": [
                {
                    "date_range": [None, None],
                    "column_map": col_map,
                }
            ],
        }
        results = transform_nrs_forenames(archived, converted, cfg)
        assert len(results) == 1

        with open(results[0]) as f:
            rows = list(csv.DictReader(f))
        assert rows[0]["sex"] == "M"  # Boy -> M
        assert rows[1]["sex"] == "F"  # Girl -> F
        assert rows[0]["name"] == "Jack"


class TestNRSSurnames:
    def test_basic(self, tmp_path):
        import openpyxl

        from pipeline.transforms.gb_scotland import transform_nrs_surnames

        archived = tmp_path / "archived"
        converted = tmp_path / "converted"
        archived.mkdir()
        converted.mkdir()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"
        # skip_rows=1, so row 0 is skipped metadata
        ws.append(["Metadata line"])
        # Header row (row index 1 after skip)
        ws.append(["Year", "Surname", "Number", "Rank"])
        ws.append([2020, "Smith", 500, 1])
        ws.append([2020, "Brown", 400, 2])
        path = archived / "nrs_surnames.xlsx"
        wb.save(path)
        wb.close()

        cfg = {
            "store": {},
            "schema_versions": [
                {"date_range": [None, None], "sheet_name": "Data", "skip_rows": 1}
            ],
        }
        results = transform_nrs_surnames(archived, converted, cfg)
        assert len(results) == 1

        with open(results[0]) as f:
            rows = list(csv.DictReader(f))
        assert len(rows) == 2
        assert rows[0]["name"] == "Smith"
        assert rows[0]["count"] == "500"
        assert rows[0]["year"] == "2020"


# ---------------------------------------------------------------------------
# England & Wales (ONS) transforms
# ---------------------------------------------------------------------------


class TestONSForenames:
    def test_basic(self, tmp_path):
        import openpyxl

        from pipeline.transforms.gb_england_wales import transform_ons_forenames

        archived = tmp_path / "archived"
        converted = tmp_path / "converted"
        archived.mkdir()
        converted.mkdir()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Table_2"
        # 4 metadata rows to skip (default skip_rows=4)
        for _ in range(4):
            ws.append(["metadata"])
        # Header
        ws.append(["Rank", "Name", "Count"])
        # Data
        ws.append([1, "Oliver", 5000])
        ws.append([2, "George", 4000])
        path = archived / "boys2020.xlsx"
        wb.save(path)
        wb.close()

        cfg = {
            "store": {"filename_date_re": r"(\d{4})"},
            "schema_versions": [
                {"date_range": [None, None], "sheet_candidates": ["Table_2"], "skip_rows": 4}
            ],
        }
        results = transform_ons_forenames(archived, converted, cfg)
        assert len(results) == 1

        with open(results[0]) as f:
            rows = list(csv.DictReader(f))
        assert rows[0]["name"] == "Oliver"
        assert rows[0]["sex"] == "M"  # "boy" in filename
        assert rows[0]["year"] == "2020"

    def test_notes_before_header(self, tmp_path):
        import openpyxl

        from pipeline.transforms.gb_england_wales import transform_ons_forenames

        archived = tmp_path / "archived"
        converted = tmp_path / "converted"
        archived.mkdir()
        converted.mkdir()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Table_2"
        # skip_rows=2 metadata
        ws.append(["metadata"])
        ws.append(["metadata"])
        # Notes before header
        ws.append(["Note: some explanation", None, None])
        # Header row (detected by "Rank" in cell0)
        ws.append(["Rank", "Name", "Count"])
        # Data
        ws.append([1, "Amelia", 3000])
        path = archived / "girls2021.xlsx"
        wb.save(path)
        wb.close()

        cfg = {
            "store": {"filename_date_re": r"(\d{4})"},
            "schema_versions": [
                {
                    "date_range": [None, None],
                    "sheet_candidates": ["Table_2"],
                    "skip_rows": 2,
                    "notes_before_header": True,
                }
            ],
        }
        results = transform_ons_forenames(archived, converted, cfg)
        assert len(results) == 1

        with open(results[0]) as f:
            rows = list(csv.DictReader(f))
        assert rows[0]["name"] == "Amelia"
        assert rows[0]["sex"] == "F"  # "girl" in filename


class TestEnglandSurnames:
    def test_basic(self, tmp_path):
        from pipeline.transforms.gb_england_wales import transform_england_surnames

        archived = tmp_path / "archived"
        converted = tmp_path / "converted"
        archived.mkdir()
        converted.mkdir()

        (archived / "surnames.csv").write_text("Surname,Count\nSmith,1000\nJones,800\n")

        cfg = {
            "store": {},
            "schema_versions": [
                {"date_range": [None, None], "columns": {"name": "Surname", "count": "Count"}}
            ],
        }
        results = transform_england_surnames(archived, converted, cfg)
        assert len(results) == 1

        with open(results[0]) as f:
            rows = list(csv.DictReader(f))
        assert rows[0]["name"] == "Smith"
        assert rows[0]["count"] == "1000"


# ---------------------------------------------------------------------------
# Northern Ireland (NISRA) transforms
# ---------------------------------------------------------------------------


class TestNISRAForenames:
    def test_basic(self, tmp_path):
        import openpyxl

        from pipeline.transforms.gb_northern_ireland import transform_nisra_forenames

        archived = tmp_path / "archived"
        converted = tmp_path / "converted"
        archived.mkdir()
        converted.mkdir()

        wb = openpyxl.Workbook()
        # Boys sheet
        ws = wb.active
        ws.title = "Boys"
        # 2 metadata rows to skip
        ws.append(["metadata"])
        ws.append(["metadata"])
        # Header with year columns (each year = 3 cols: "{YYYY} Name", count, rank)
        ws.append(
            [
                "2020 Name",
                "Number of Babies",
                "Rank",
                "2021 Name",
                "Number of Babies",
                "Rank",
            ]
        )
        # Data
        ws.append(["Jack", 100, 1, "Jack", 110, 1])
        ws.append(["Noah", 90, 2, "Noah", "..", ".."])  # 2021 suppressed

        # Girls sheet
        ws2 = wb.create_sheet("Girls")
        ws2.append(["metadata"])
        ws2.append(["metadata"])
        ws2.append(["2020 Name", "Number of Babies", "Rank"])
        ws2.append(["Emily", 80, 1])

        path = archived / "nisra_2020.xlsx"
        wb.save(path)
        wb.close()

        cfg = {
            "store": {"filename_date_re": r"(\d{4})"},
            "schema_versions": [
                {
                    "date_range": [None, None],
                    "sheets": [
                        {"sheet_name": "Boys", "sex": "M"},
                        {"sheet_name": "Girls", "sex": "F"},
                    ],
                    "skip_rows": 2,
                }
            ],
        }
        results = transform_nisra_forenames(archived, converted, cfg)
        assert len(results) == 1

        with open(results[0]) as f:
            rows = list(csv.DictReader(f))

        # Jack 2020, Jack 2021, Noah 2020 (Noah 2021 suppressed), Emily 2020
        names_years = [(r["name"], r["year"], r["sex"]) for r in rows]
        assert ("Jack", "2020", "M") in names_years
        assert ("Jack", "2021", "M") in names_years
        assert ("Noah", "2020", "M") in names_years
        assert ("Emily", "2020", "F") in names_years
        # Noah 2021 should be filtered out (suppressed "..")
        noah_2021 = [r for r in rows if r["name"] == "Noah" and r["year"] == "2021"]
        assert len(noah_2021) == 0


# ---------------------------------------------------------------------------
# Ireland (CSO) transforms
# ---------------------------------------------------------------------------


class TestCSOForenames:
    def test_basic(self, tmp_path):
        from pipeline.transforms.ie import transform_cso_boys

        archived = tmp_path / "archived"
        converted = tmp_path / "converted"
        archived.mkdir()
        converted.mkdir()

        # PxStat CSV with BOM
        content = "\ufeffStatistic Label,Year,Boys Names,UNIT,VALUE\n"
        content += "Births,2020,Jack,Number,500\n"
        content += "Births,2020,James,Number,bad_value\n"
        (archived / "boys.csv").write_text(content, encoding="utf-8")

        cfg = {
            "store": {},
            "schema_versions": [
                {"date_range": [None, None], "name_column": "Boys Names", "sex": "M"}
            ],
        }
        results = transform_cso_boys(archived, converted, cfg)
        assert len(results) == 1

        with open(results[0]) as f:
            rows = list(csv.DictReader(f))
        assert rows[0]["name"] == "Jack"
        assert rows[0]["count"] == "500"
        assert rows[0]["sex"] == "M"
        # Bad value passed through for rejection detection
        assert rows[1]["count"] == "bad_value"


class TestCSOSurnames:
    def test_rank_meaning(self, tmp_path):
        from pipeline.transforms.ie import transform_cso_surnames

        archived = tmp_path / "archived"
        converted = tmp_path / "converted"
        archived.mkdir()
        converted.mkdir()

        content = "\ufeffStatistic Label,Year,Surnames,UNIT,VALUE\n"
        content += "Rank,2020,Murphy,Number,1\n"
        (archived / "surnames.csv").write_text(content, encoding="utf-8")

        cfg = {
            "store": {},
            "schema_versions": [
                {"date_range": [None, None], "name_column": "Surnames", "value_meaning": "rank"}
            ],
        }
        results = transform_cso_surnames(archived, converted, cfg)
        with open(results[0]) as f:
            rows = list(csv.DictReader(f))
        # VALUE=1 goes to rank column, count is empty
        assert rows[0]["rank"] == "1"
        assert rows[0]["count"] == ""

    def test_count_meaning(self, tmp_path):
        from pipeline.transforms.ie import transform_cso_surnames

        archived = tmp_path / "archived"
        converted = tmp_path / "converted"
        archived.mkdir()
        converted.mkdir()

        content = "\ufeffStatistic Label,Year,Surnames,UNIT,VALUE\n"
        content += "Count,2020,Murphy,Number,5000\n"
        (archived / "surnames.csv").write_text(content, encoding="utf-8")

        cfg = {
            "store": {},
            "schema_versions": [
                {"date_range": [None, None], "name_column": "Surnames", "value_meaning": "count"}
            ],
        }
        results = transform_cso_surnames(archived, converted, cfg)
        with open(results[0]) as f:
            rows = list(csv.DictReader(f))
        # VALUE=5000 goes to count column, rank is empty
        assert rows[0]["count"] == "5000"
        assert rows[0]["rank"] == ""
